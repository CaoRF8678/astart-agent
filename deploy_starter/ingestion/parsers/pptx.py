import io
import zipfile

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from ingestion.errors import (
    DocumentParseError,
)
from ingestion.parsers.base import (
    DocumentParser,
)
from ingestion.text_processing import (
    merge_short_blocks,
    normalize_text,
    split_long_text,
)
from ingestion.types import ParsedBlock


class PPTXParser(DocumentParser):

    def parse(
        self,
        data: bytes,
    ) -> list[ParsedBlock]:
        try:
            presentation = Presentation(
                io.BytesIO(data)
            )
        except (
            PackageNotFoundError,
            zipfile.BadZipFile,
            ValueError,
            KeyError,
        ) as exc:
            raise DocumentParseError(
                "无法解析 PPTX 文档。"
            ) from exc

        parsed: list[ParsedBlock] = []

        for slide_index, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            title_shape = slide.shapes.title

            title = (
                normalize_text(
                    title_shape.text
                )
                if title_shape is not None
                else ""
            )

            shapes = sorted(
                slide.shapes,
                key=lambda shape: (
                    shape.top,
                    shape.left,
                ),
            )

            units: list[str] = []

            if title:
                units.append(title)

            for shape in shapes:
                if (
                    title_shape is not None
                    and shape.shape_id
                    == title_shape.shape_id
                ):
                    continue

                if getattr(
                    shape,
                    "has_table",
                    False,
                ):
                    rows: list[str] = []

                    for row in shape.table.rows:
                        cells = [
                            normalize_text(
                                cell.text
                            )
                            for cell in row.cells
                        ]

                        cells = [
                            cell
                            for cell in cells
                            if cell
                        ]

                        if cells:
                            rows.append(
                                " | ".join(cells)
                            )

                    if rows:
                        units.append(
                            "\n".join(rows)
                        )

                    continue

                if getattr(
                    shape,
                    "has_text_frame",
                    False,
                ):
                    text = normalize_text(
                        shape.text
                    )

                    if text:
                        units.append(text)

            atomic: list[str] = []

            for unit in units:
                atomic.extend(
                    split_long_text(unit)
                )

            for content in merge_short_blocks(
                atomic
            ):
                locator: dict = {
                    "slide_start": slide_index,
                    "slide_end": slide_index,
                }

                if title:
                    locator["title"] = title

                block = self._make_block(
                    content=content,
                    locator=locator,
                )

                if block is not None:
                    parsed.append(block)

        if not parsed:
            raise DocumentParseError(
                "PPTX 中没有提取到有效文本。"
            )

        return parsed